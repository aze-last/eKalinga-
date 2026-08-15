import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def remove_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

def add_callout(doc, title, text, callout_type="policy"):
    # policy: green border/tint, tip: amber border/tint, note: blue border/tint
    if callout_type == "policy":
        bg_color = "F0FDF4"
        border_color = "15803D"
        icon = "🛡️ SYSTEM POLICY & INVARIANT"
    elif callout_type == "tip":
        bg_color = "FFFBEB"
        border_color = "D97706"
        icon = "💡 OPERATIONAL PRO-TIP"
    else:
        bg_color = "F8FAFC"
        border_color = "475569"
        icon = "📌 NOTE"

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.8)
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="28" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"{icon}: ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor.from_string(border_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9)
    r_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_section_header_card(doc, num, title, domain_tag):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.4)
    table.columns[1].width = Inches(1.4)
    
    c0 = table.cell(0, 0)
    c1 = table.cell(0, 1)
    
    set_cell_background(c0, "064E3B") # Dark Emerald
    set_cell_background(c1, "064E3B")
    set_cell_margins(c0, top=160, bottom=160, left=180, right=80)
    set_cell_margins(c1, top=160, bottom=160, left=80, right=180)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r_sub = p0.add_run(f"SECTION {num}  •  ")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(0x6E, 0xE7, 0xB7) # Light Mint
    r_sub.bold = True
    
    r_main = p0.add_run(title)
    r_main.font.name = "Segoe UI"
    r_main.font.size = Pt(13)
    r_main.bold = True
    r_main.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(0)
    r_tag = p1.add_run(f"[{domain_tag}]")
    r_tag.font.name = "Segoe UI"
    r_tag.font.size = Pt(8.5)
    r_tag.bold = True
    r_tag.font.color.rgb = RGBColor(0xFC, 0xD3, 0x4D) # Amber Gold
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def build_manual():
    base_dir = r"c:\Users\ASUS\source\repos\eKalinga-"
    screenshots_dir = os.path.join(base_dir, "docs", "manual_screenshots")
    docx_out = os.path.join(base_dir, "docs", "User-Manual-eKalingaPlus.docx")
    docx_ayuda = os.path.join(base_dir, "docs", "User-Manual-Ayuda.docx")
    md_out = os.path.join(base_dir, "docs", "User-Manual-eKalingaPlus.md")

    sections_data = [
        {
            "num": "01",
            "title": "System Architecture & Login Authentication",
            "domain": "SYSTEM SETUP",
            "img": "01_login_window.png",
            "summary": "The eKalinga+ application launches with a dual-branded, security-hardened authentication portal linking the official LGU Seal and the eKalinga+ operational core.",
            "workflow": [
                "1. **Identity & Connectivity Display:** The left panel surfaces the Province of Davao del Sur / Municipality of Sulop seal, system serial, and active database connection indicator.",
                "2. **Credentials Entry:** Enter your assigned username/email and password. The system supports full keyboard navigation (Press Enter to execute sign in).",
                "3. **Demo & Default Credentials:** For local offline deployment, pre-configured roles are available (`admin@barangay.local` / `admin123`).",
                "4. **OTP Multi-Factor Security:** Operators can toggle email OTP verification for enhanced authorization during staff shifts."
            ],
            "rule": "Only authorized staff accounts may sign in. Password lockouts occur after consecutive failed attempts.",
            "tip": "Keep the Local MySQL service running in XAMPP or Windows Services before launching the application."
        },
        {
            "num": "02",
            "title": "Database Connection & Offline Preset Configuration",
            "domain": "SYSTEM SETUP",
            "img": "02_connection_settings.png",
            "summary": "eKalinga+ features seamless multi-tier database synchronisation across Local MySQL, LAN Server, and Remote Cloud tiers with auto-migration.",
            "workflow": [
                "1. **Connection Presets:** Switch between `Local (127.0.0.1:3306)`, `Network (LAN)`, or `Remote Cloud` directly from the login settings window.",
                "2. **Live Connectivity Test:** Click 'Test Connection' to verify socket reachability, schema integrity, and response latency before committing changes.",
                "3. **Zero-Configuration Bootstrap:** Missing database tables and migrations are automatically created by the `StartupMigrationCoordinator` upon initial startup."
            ],
            "rule": "Standard workstations must keep the active preset on 'Local' unless explicitly directed by the IT Administrator.",
            "tip": "Click 'Test Connection' before saving credentials to prevent authentication lockouts."
        },
        {
            "num": "03",
            "title": "Executive Dashboard & Operational Command Center",
            "domain": "OPERATIONS",
            "img": "03_dashboard_page.png",
            "summary": "The centralized landing hub providing real-time KPI metrics, assistance totals, remaining budget capacity, and one-click navigation to all active modules.",
            "workflow": [
                "1. **Summary Analytics:** Review total registered beneficiaries (40,561), civil registry links, active assistance cases, and remaining funding pools.",
                "2. **Direct Module Dispatch:** Click any module tile (Masterlist, Aid Requests, Budget Waterfall, Project Distribution, Cash-for-Work, Seminar Attendance, GGMS Transactions, Reports) to open its operational workspace.",
                "3. **Background Sync Status:** Bottom indicator shows real-time online/offline and GGMS central server synchronization states."
            ],
            "rule": "The Dashboard layout is locked to maintain interface consistency across all municipal workstations.",
            "tip": "Use the sidebar navigation drawer for rapid multitasking across different active distribution events."
        },
        {
            "num": "04",
            "title": "Validated Beneficiaries Masterlist & Digital ID System",
            "domain": "OPERATIONS",
            "img": "04_masterlist_page.png",
            "summary": "The master registry of all validated citizens in the municipality, featuring comprehensive profiling, demographic tagging, and QR/Barcode digital ID generation.",
            "workflow": [
                "1. **Real-time Search & Filter:** Filter by demographics (Senior Citizens, PWDs, 4Ps, Solo Parents, Youth, Farmers, Fisherfolk).",
                "2. **Photo Capture & Upload:** Directly capture beneficiary portrait photos via attached USB PC Camera or upload image files with integrated square cropping.",
                "3. **Digital eKard Printing:** Select a beneficiary and click 'Print Digital ID' to generate compliant QR/Barcode badges for offline rapid scanning during distribution events.",
                "4. **CRS Municipal Sync:** Synchronize citizen validation status directly against the central Civil Registry System."
            ],
            "rule": "Never create duplicate entries. Use the CRS Sync ID or National ID number for deduplication.",
            "tip": "Batch print QR digital badges prior to large distribution events to maximize onsite throughput."
        },
        {
            "num": "05",
            "title": "Aid Requests & Assistance Case State Machine",
            "domain": "FINANCIAL",
            "img": "05_aid_request_page.png",
            "summary": "Manages individual walk-in assistance claims (Medical, Burial, Educational, Emergency Food) following a strict 4-stage approval and release pipeline.",
            "workflow": [
                "1. **State Pipeline:** Requests progress strictly through: `Pending` ➔ `Under Review` ➔ `Approved` ➔ `Released` (or `Rejected`/`Cancelled`).",
                "2. **New Request Filing:** Search beneficiary from masterlist, select Assistance Category, enter requested amount, and attach required proof documents.",
                "3. **Approval & Allocation:** Admin reviews the case and sets the official `Approved Amount`. Funds cannot be disbursed without prior approval.",
                "4. **Fund Release & Budget Waterfall Deduction:** Releasing assistance automatically creates a `BudgetLedgerEntry` that consumes from the earmarked Assistance Case pool, maintaining strict fiscal accountability."
            ],
            "rule": "Every released assistance case must be coupled with an explicit ApprovedAmount and authorized officer.",
            "tip": "Attach verified digital copies of hospital bills or death certificates to accelerate the approval stage."
        },
        {
            "num": "06",
            "title": "Budget Management & Dual-Stream Waterfall Ledger",
            "domain": "FINANCIAL",
            "img": "06_budget_page.png",
            "summary": "The financial engine of eKalinga+, orchestrating government appropriations (GGMS) and private donations through an automated waterfall disbursement hierarchy.",
            "workflow": [
                "1. **Dual Inflow Streams:** Tracks GGMS Government Allocations alongside Private Donations with reference numbers and proof receipts.",
                "2. **Budget Earmarking:** Funds are categorized into sub-budgets (Aid Requests, Cash-for-Work, Project Distributions, Seminars).",
                "3. **Waterfall Hierarchy:** Releases automatically deduct first from the specific earmarked project fund, then cascade to the general assistance pool.",
                "4. **Create New Project (1:1 Funding):** Directly spawn projects from a selected donation/GGMS fund without breaking ledger continuity."
            ],
            "rule": "Independent or disconnected budget records that bypass the waterfall ledger are strictly forbidden.",
            "tip": "Audit remaining project budgets weekly to ensure adequate reserve for incoming emergency claims."
        },
        {
            "num": "07",
            "title": "Project Distribution & Disbursement Execution",
            "domain": "FINANCIAL",
            "img": "07_project_distribution_page.png",
            "summary": "High-throughput operational module for disbursing bulk goods and financial assistance to enrolled masterlist beneficiaries.",
            "workflow": [
                "1. **Project Selection:** Load active distribution project (Cash / In-Kind Goods) linked to its funding source.",
                "2. **Participant Queue:** Displays full participant roster categorized into `RELEASED / CLAIMED`, `PENDING (Requirements Lacking)`, and `UNRELEASED`.",
                "3. **Rapid Identification:** Scan physical QR ID card, type Beneficiary ID, or double-click list item.",
                "4. **Requirements Verification:** Check mandatory attachments (Cedula, Barangay Certificate). Click 'Confirm Release' to payout or 'Mark Pending' if documents are missing."
            ],
            "rule": "Beneficiary enrollment is manual or batch-selected from approved masterlist records only (no demographic auto-enrollment).",
            "tip": "Use a 2D barcode scanner gun set to USB HID mode for sub-second verification per beneficiary."
        },
        {
            "num": "08",
            "title": "Cash-for-Work Attendance & Payout Operations",
            "domain": "FINANCIAL",
            "img": "08_cash_for_work_page.png",
            "summary": "Complete lifecycle management for Cash-for-Work community projects, daily worker attendance logging, and automated wage payout disbursement.",
            "workflow": [
                "1. **Event Workspace:** Create work activity (e.g. Barangay Clean-up, Tree Planting), assign daily wage rate and scheduled days.",
                "2. **Beneficiary Enrollment:** Add registered workers from the masterlist to the event roster.",
                "3. **Barcode/OCR Attendance Scan:** Capture morning time-in and afternoon time-out using USB barcode gun or live PC camera OCR.",
                "4. **Automated Payroll Calculation:** System aggregates verified attendance logs into the Payout view (Calculates Days Rendered × Daily Wage Rate)."
            ],
            "rule": "Payouts require verified attendance records and pull directly from the CashForWork budget bucket.",
            "tip": "Ensure camera lighting is adequate when utilizing the built-in OCR visual scanner."
        },
        {
            "num": "09",
            "title": "Cash-for-Work Payout Disbursement Ledger",
            "domain": "FINANCIAL",
            "img": "09_cash_for_work_payout_page.png",
            "summary": "The payroll disbursement terminal for releasing wages to Cash-for-Work participants upon event completion.",
            "workflow": [
                "1. **Review Earned Wages:** System computes exact payout per participant based on logged attendance days.",
                "2. **Disbursement Confirmation:** Process individual or batch payouts with receipt generation.",
                "3. **Audit Trail:** Every released wage logs a ledger entry tied to the Cash-for-Work budget stream."
            ],
            "rule": "Ensure all daily attendance logs are finalized before initiating payout disbursement.",
            "tip": "Print the disbursement liquidation summary immediately after completing payroll."
        },
        {
            "num": "10",
            "title": "Seminar & Training Attendance Verification",
            "domain": "OPERATIONS",
            "img": "10_seminar_attendance_page.png",
            "summary": "Specialized module for conducting training sessions, workshops, and livelihood seminars with real-time barcode/QR scan tracking.",
            "workflow": [
                "1. **Seminar Setup:** Define seminar topic, venue, date, and target participants.",
                "2. **Rapid Scanner Entry:** Position camera or handheld scanner to log attendees as they arrive.",
                "3. **Instant Attendance Summary:** Real-time stats surface total attendees, check-in timestamps, and demographic breakdown."
            ],
            "rule": "All seminar projects originate from an earmarked funding stream in the Budget module.",
            "tip": "Keep the live log window open to verify successful scans as attendees enter the venue."
        },
        {
            "num": "11",
            "title": "GGMS Consolidated Transactions & Sync Audit",
            "domain": "OPERATIONS",
            "img": "11_ggms_transactions_page.png",
            "summary": "Centralized transaction monitoring linking local eKalinga+ disbursements with the municipal Government Grants Management System (GGMS).",
            "workflow": [
                "1. **Synced Record Inspection:** Review all assistance claims, distributions, and payouts transmitted to GGMS.",
                "2. **Filter by Sync State:** Filter by Synchronized, Pending Sync, or Conflict.",
                "3. **Reconciliation Audit:** Inspect timestamp, batch reference, transaction hash, and approving officer details."
            ],
            "rule": "Ensure periodic sync is active when operating on connected LAN/Remote networks.",
            "tip": "Check the conflict tab after reconnecting an offline workstation to resolve duplicate claims."
        },
        {
            "num": "12",
            "title": "Reports, Analytics & Official Export Engine",
            "domain": "GOVERNANCE",
            "img": "12_reports_page.png",
            "summary": "Comprehensive document generation engine creating formatted government reports, audit sheets, and statistical summaries.",
            "workflow": [
                "1. **Select Report Template:** Masterlist Summary, Distribution Liquidation Sheet, Aid Assistance Ledger, or Cash-for-Work Payroll.",
                "2. **Date & Category Filtering:** Define report timeframes and program categories.",
                "3. **Print Preview & Export:** Preview formatted documents with official LGU headers and export directly to PDF or Excel."
            ],
            "rule": "All official exported reports automatically include the municipal watermark and signature fields.",
            "tip": "Use PDF export for archived official submissions and Excel for statistical aggregation."
        },
        {
            "num": "13",
            "title": "Scanning Portal & Remote Scanner Gateway",
            "domain": "OPERATIONS",
            "img": "13_scanning_portal_page.png",
            "summary": "Network terminal for connecting mobile devices and secondary hardware scanners across the local barangay hall network.",
            "workflow": [
                "1. **Scanner Session PIN:** Generate secure, time-expiring PIN codes to pair remote Android/iOS scanning terminals.",
                "2. **Live Feed Monitoring:** Monitor multiple active scanner gates simultaneously during massive distribution events.",
                "3. **Hardware Scanner Integration:** Compatible with USB 1D/2D Barcode Guns, PC Webcams, and Serial scanners without extra drivers."
            ],
            "rule": "Session PINs automatically expire after the configured timeout period for security.",
            "tip": "Assign separate scanner PINs to each gate operator during municipal-wide relief operations."
        },
        {
            "num": "14",
            "title": "System Settings, Security & Database Backup Recovery",
            "domain": "SECURITY",
            "img": "14_settings_window.png",
            "summary": "Central administration console governing system identity, SuperAdmin user management, password-protected overlays, and database backup chains.",
            "workflow": [
                "1. **System Profile:** Set municipal name, address, branding logo, and theme preferences.",
                "2. **User Management & Granular Permissions:** SuperAdmin manages staff accounts and assigns granular module access checkboxes.",
                "3. **Password Protection:** Sensitive areas (App Database, Remote Snapshot, GGMS Source) require admin password re-entry.",
                "4. **Database Backup & Recovery:** Execute Full Baseline Backups, Incremental Backups, and point-in-time Restores directly from the UI."
            ],
            "rule": "Never delete database rows. Only soft deletions (`IsDeleted`) are permitted. Always create a Full Backup before major server migrations.",
            "tip": "Schedule incremental daily backups to an external encrypted USB drive for disaster recovery."
        }
    ]

    doc = docx.Document()

    # Page Margins
    for s_item in doc.sections:
        s_item.top_margin = Inches(0.7)
        s_item.bottom_margin = Inches(0.7)
        s_item.left_margin = Inches(0.75)
        s_item.right_margin = Inches(0.75)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # ══════════════════════════════════════════════════════
    # 1. EXECUTIVE COVER PAGE
    # ══════════════════════════════════════════════════════
    # Top Decorative Emerald Banner Table
    top_band = doc.add_table(rows=1, cols=1)
    top_band.alignment = WD_TABLE_ALIGNMENT.CENTER
    top_band.autofit = False
    top_band.columns[0].width = Inches(7.0)
    c_band = top_band.cell(0, 0)
    set_cell_background(c_band, "064E3B") # Deep Forest Emerald
    set_cell_margins(c_band, top=140, bottom=140, left=180, right=180)
    remove_cell_borders(c_band)
    
    p_band = c_band.paragraphs[0]
    p_band.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_band = p_band.add_run("REPUBLIC OF THE PHILIPPINES  •  PROVINCE OF DAVAO DEL SUR  •  MUNICIPALITY OF SULOP")
    r_band.font.name = "Segoe UI"
    r_band.font.size = Pt(8)
    r_band.bold = True
    r_band.font.color.rgb = RGBColor(0xA7, 0xF3, 0xD0) # Mint Accent

    # Spacer
    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(16)
    p_sp1.paragraph_format.space_after = Pt(0)

    # Dual Logos Table (LGU Seal + eKalinga+ Logo)
    logo_table = doc.add_table(rows=1, cols=2)
    logo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    logo_table.autofit = False
    logo_table.columns[0].width = Inches(3.5)
    logo_table.columns[1].width = Inches(3.5)
    
    c_seal = logo_table.cell(0, 0)
    c_logo = logo_table.cell(0, 1)
    remove_cell_borders(c_seal)
    remove_cell_borders(c_logo)
    
    seal_path = os.path.join(base_dir, "Images", "7a9e9592-b36f-4983-afcb-d9dea4caa022.jpg")
    ekalinga_path = os.path.join(base_dir, "Images", "Gemini_Generated_Image_1ivs1t1ivs1t1ivs-removebg-preview.png")
    
    p_seal = c_seal.paragraphs[0]
    p_seal.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(seal_path):
        p_seal.add_run().add_picture(seal_path, width=Inches(1.3))
        
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(ekalinga_path):
        p_logo.add_run().add_picture(ekalinga_path, width=Inches(1.5))

    # Main Titles
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(2)
    
    r_main_title = p_title.add_run("eKalinga+\n")
    r_main_title.font.name = "Segoe UI"
    r_main_title.font.size = Pt(32)
    r_main_title.bold = True
    r_main_title.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)
    
    r_sub1 = p_title.add_run("Ayuda Management & Unified Distribution System\n")
    r_sub1.font.name = "Segoe UI"
    r_sub1.font.size = Pt(16)
    r_sub1.bold = True
    r_sub1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r_sub2 = p_title.add_run("Comprehensive Operational & Technical Flow User Manual")
    r_sub2.font.name = "Segoe UI"
    r_sub2.font.size = Pt(11.5)
    r_sub2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Document Control Metadata Box
    doc_card = doc.add_table(rows=4, cols=2)
    doc_card.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc_card.autofit = False
    doc_card.columns[0].width = Inches(2.2)
    doc_card.columns[1].width = Inches(4.6)
    
    meta_entries = [
        ("System Release Version", "v1.0.6 (Desktop Enterprise Edition)"),
        ("Implementing Authority", "Municipality of Sulop / Barangay Social Welfare"),
        ("Database Classification", "Multi-Tier Dual Ledger (Local / LAN / Cloud Sync)"),
        ("Official Target Audience", "System Administrators, Encoders, Field Gate Officers")
    ]
    
    for idx, (label, val) in enumerate(meta_entries):
        row = doc_card.rows[idx]
        c0 = row.cells[0]
        c1 = row.cells[1]
        bg = "F8FAFC" if idx % 2 == 0 else "F1F5F9"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(8.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(8.5)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)

    # Bottom Document Subtitle
    p_bot = doc.add_paragraph()
    p_bot.paragraph_format.space_before = Pt(28)
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = p_bot.add_run("CONFIDENTIAL & PROPRIETARY  •  OFFICIAL LGU WORKSTATION USE ONLY")
    r_bot.font.name = "Segoe UI"
    r_bot.font.size = Pt(8)
    r_bot.bold = True
    r_bot.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════
    # 2. TABLE OF CONTENTS
    # ══════════════════════════════════════════════════════
    p_toc_head = doc.add_paragraph()
    p_toc_head.paragraph_format.space_before = Pt(6)
    p_toc_head.paragraph_format.space_after = Pt(2)
    r_toc = p_toc_head.add_run("Table of Contents")
    r_toc.font.name = "Segoe UI"
    r_toc.font.size = Pt(18)
    r_toc.bold = True
    r_toc.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)

    p_toc_sub = doc.add_paragraph()
    p_toc_sub.paragraph_format.space_after = Pt(10)
    r_toc_sub = p_toc_sub.add_run("Operational Directory & Module Navigation Matrix")
    r_toc_sub.font.name = "Segoe UI"
    r_toc_sub.font.size = Pt(9.5)
    r_toc_sub.italic = True
    r_toc_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # TOC Table
    toc_table = doc.add_table(rows=len(sections_data) + 1, cols=3)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_table.autofit = False

    col_widths = [Inches(0.7), Inches(4.7), Inches(1.4)]
    for i, col in enumerate(toc_table.columns):
        col.width = col_widths[i]

    # Header Row
    headers = ["Sec #", "Operational Module & Procedure", "Domain Tag"]
    hdr_cells = toc_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "064E3B")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=90, right=90)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data Rows
    for idx, s in enumerate(sections_data):
        row_cells = toc_table.rows[idx + 1].cells
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i in range(3):
            row_cells[i].width = col_widths[i]
            set_cell_background(row_cells[i], bg)
            set_cell_margins(row_cells[i], top=65, bottom=65, left=90, right=90)

        # Sec #
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(s['num'])
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(8.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)

        # Title
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(s['title'])
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(8.5)
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        # Scope Tag
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(s['domain'])
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(7.5)
        r2.bold = True
        if s['domain'] == "SYSTEM SETUP":
            r2.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue
        elif s['domain'] == "FINANCIAL":
            r2.font.color.rgb = RGBColor(0xD9, 0x77, 0x06) # Amber
        elif s['domain'] == "SECURITY":
            r2.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED) # Purple
        else:
            r2.font.color.rgb = RGBColor(0x05, 0x96, 0x69) # Green

    doc.add_page_break()

    # ══════════════════════════════════════════════════════
    # 3. MANUAL SECTIONS
    # ══════════════════════════════════════════════════════

    md_lines = [
        "# eKalinga+ / Ayuda Management System",
        "## Comprehensive Operations & Flow User Manual",
        "**Official Desktop Edition (v1.0.6)**  •  *Municipality of Sulop, Province of Davao del Sur*",
        "\n---\n",
        "## Table of Contents\n"
    ]

    for s in sections_data:
        md_lines.append(f"- [{s['num']}. {s['title']}](#{s['num']}-{s['title'].lower().replace(' ', '-').replace('&', '').replace('/', '').replace(',', '')})")

    md_lines.append("\n---\n")

    for s in sections_data:
        # Styled Header Card
        add_section_header_card(doc, s['num'], s['title'], s['domain'])

        # Overview Paragraph
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(2)
        p_desc.paragraph_format.space_after = Pt(6)
        r_desc = p_desc.add_run(s['summary'])
        r_desc.font.size = Pt(9.5)

        # Image Frame Box
        img_file = os.path.join(screenshots_dir, s['img'])
        if os.path.exists(img_file):
            try:
                img_table = doc.add_table(rows=1, cols=1)
                img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                img_table.autofit = False
                img_table.columns[0].width = Inches(6.8)
                
                c_img = img_table.cell(0, 0)
                set_cell_background(c_img, "F8FAFC")
                set_cell_margins(c_img, top=100, bottom=100, left=100, right=100)
                
                # Light subtle border
                tcPr = c_img._tc.get_or_add_tcPr()
                borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
                tcPr.append(borders)
                
                p_in_img = c_img.paragraphs[0]
                p_in_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_in_img.paragraph_format.space_before = Pt(2)
                p_in_img.paragraph_format.space_after = Pt(2)
                p_in_img.add_run().add_picture(img_file, width=Inches(6.4))

                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(8)
                r_cap = p_cap.add_run(f"Figure {s['num']}.1: {s['title']} — Operational View")
                r_cap.font.name = "Segoe UI"
                r_cap.font.size = Pt(8)
                r_cap.italic = True
                r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            except Exception as e:
                print(f"Failed to add image {img_file}: {e}")

        # Workflow steps header
        p_wf_head = doc.add_paragraph()
        p_wf_head.paragraph_format.space_before = Pt(4)
        p_wf_head.paragraph_format.space_after = Pt(4)
        r_wf_head = p_wf_head.add_run("📋 Operational Step-by-Step Workflow:")
        r_wf_head.bold = True
        r_wf_head.font.size = Pt(10)
        r_wf_head.font.color.rgb = RGBColor(0x06, 0x4E, 0x3B)

        for step in s['workflow']:
            p_step = doc.add_paragraph()
            p_step.paragraph_format.left_indent = Inches(0.15)
            p_step.paragraph_format.space_after = Pt(2.5)
            parts = step.split("**")
            is_bold = False
            for part in parts:
                r_step = p_step.add_run(part)
                r_step.font.size = Pt(9)
                if is_bold:
                    r_step.bold = True
                    r_step.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                else:
                    r_step.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                is_bold = not is_bold

        # Add Callouts
        add_callout(doc, "System Policy", s['rule'], callout_type="policy")
        add_callout(doc, "Operator Tip", s['tip'], callout_type="tip")

        # Markdown lines
        md_lines.append(f"## {s['num']}. {s['title']} `[{s['domain']}]`\n")
        md_lines.append(f"{s['summary']}\n")
        md_lines.append(f"![Figure {s['num']}.1: {s['title']}](manual_screenshots/{s['img']})\n")
        md_lines.append("### Operational Workflow & Step-by-Step Flow:\n")
        for step in s['workflow']:
            md_lines.append(f"- {step}")
        md_lines.append(f"\n> **🛡️ System Policy:** {s['rule']}\n")
        md_lines.append(f"> **💡 Pro-Tip:** {s['tip']}\n\n---\n")

    # Save to User-Manual-eKalingaPlus.docx and User-Manual-Ayuda.docx
    doc.save(docx_out)
    print(f"Saved DOCX manual to: {docx_out}")

    try:
        doc.save(docx_ayuda)
        print(f"Saved DOCX manual to: {docx_ayuda}")
    except PermissionError:
        print("User-Manual-Ayuda.docx is locked in Word; saved to User-Manual-eKalingaPlus.docx.")

    with open(md_out, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    print(f"Saved Markdown manual to: {md_out}")

if __name__ == "__main__":
    build_manual()
